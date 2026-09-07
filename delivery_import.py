"""Import expected-delivery rows from PCB_BOM_Parser_Py's Excel report."""
from pathlib import Path
from decimal import Decimal, InvalidOperation
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


HEADER_ALIASES = {
    "исходное наименование": "source_name",
    "тип элемента": "component_type",
    "параметры": "parameters",
    "список на англ.": "english_name",
    "список на рус.": "russian_name",
    "наимен. произв.": "manufacturer_part_name",
    "производитель": "manufacturer",
    "количество": "quantity",
}
REQUIRED_HEADERS = {"source_name", "quantity"}
PARAMETER_FIELDS = {
    "значение": "value",
    "корпус": "case",
    "точность": "tol",
}
TYPE_ALIASES = {
    "-": "Прочее",
    "Индуктивность": "Катушка индуктивности",
}
UNIT_TRANSLATIONS = {
    "Ohm": "Ом",
    "ohm": "Ом",
    "Ω": "Ом",
    "Ω": "Ом",
    "mOhm": "мОм",
    "kOhm": "кОм",
    "KOhm": "кОм",
    "kohm": "кОм",
    "MOhm": "МОм",
    "F": "Ф",
    "mF": "мФ",
    "uF": "мкФ",
    "µF": "мкФ",
    "μF": "мкФ",
    "nF": "нФ",
    "pF": "пФ",
    "H": "Гн",
    "mH": "мГн",
    "uH": "мкГн",
    "µH": "мкГн",
    "μH": "мкГн",
    "nH": "нГн",
    "V": "В",
    "mV": "мВ",
    "uV": "мкВ",
    "A": "А",
    "mA": "мА",
    "uA": "мкА",
    "W": "Вт",
    "mW": "мВт",
    "Hz": "Гц",
    "kHz": "кГц",
    "MHz": "МГц",
    "GHz": "ГГц",
}


class DeliveryImportError(ValueError):
    pass


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _meaningful(value):
    value = _text(value)
    return "" if value == "-" else value


def _number_text(value):
    try:
        number = Decimal(value.replace(",", "."))
    except InvalidOperation:
        return value
    normalized = format(number.normalize(), "f")
    return normalized or "0"


def _translate_unit(value):
    value = value.strip()
    return UNIT_TRANSLATIONS.get(value, value)


def _parse_quantity(value, row_number):
    if isinstance(value, bool):
        raise DeliveryImportError(
            f"Строка {row_number}: количество должно быть целым положительным числом."
        )
    text = _text(value)
    if not re.fullmatch(r"[0-9]+(?:\.0+)?", text):
        raise DeliveryImportError(
            f"Строка {row_number}: количество должно быть целым положительным числом."
        )
    try:
        numeric = float(text)
        number = int(numeric)
    except (ValueError, OverflowError):
        raise DeliveryImportError(
            f"Строка {row_number}: количество должно быть целым положительным числом."
        ) from None
    if number <= 0 or number > 9007199254740991 or numeric != number:
        raise DeliveryImportError(
            f"Строка {row_number}: количество должно быть целым положительным числом."
        )
    return number


def _parse_parameters(value):
    result = {"value": "", "unit": "", "tol": "", "case": ""}
    for line in _text(value).splitlines():
        label, separator, raw_value = line.partition(":")
        if not separator:
            continue
        key = PARAMETER_FIELDS.get(label.strip().lower())
        raw_value = _meaningful(raw_value)
        if not key or not raw_value:
            continue
        if key == "value":
            match = re.match(r"^(.+?)\s+([^\s]+)$", raw_value)
            if match:
                result["value"] = _number_text(match.group(1).strip())
                result["unit"] = _translate_unit(match.group(2))
            else:
                result["value"] = _number_text(raw_value)
        elif key == "tol":
            tolerance = re.sub(r"\s+%$", "", raw_value)
            result[key] = f"{_number_text(tolerance)}%"
        else:
            result[key] = raw_value
    return result


def _limited(value, row_number, label):
    value = _meaningful(value)
    if len(value) > 255:
        raise DeliveryImportError(
            f"Строка {row_number}: поле «{label}» длиннее 255 символов."
        )
    return value


def parse_delivery_workbook(file_stream, filename):
    if Path(filename or "").suffix.lower() != ".xlsx":
        raise DeliveryImportError("Выберите отчёт в формате .xlsx.")
    try:
        workbook = load_workbook(file_stream, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError):
        raise DeliveryImportError("Не удалось прочитать файл Excel.") from None

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            raise DeliveryImportError("В отчёте нет строки заголовков.")

        columns = {}
        for index, value in enumerate(header_values):
            key = HEADER_ALIASES.get(_text(value).lower())
            if key and key not in columns:
                columns[key] = index
        missing = REQUIRED_HEADERS - set(columns)
        if missing:
            raise DeliveryImportError(
                "В отчёте должны быть столбцы «Исходное наименование» и «Количество»."
            )

        result = []
        for row_number, row in enumerate(rows, start=2):
            source_name = _text(row[columns["source_name"]]) if columns["source_name"] < len(row) else ""
            quantity_value = row[columns["quantity"]] if columns["quantity"] < len(row) else None
            if not source_name and quantity_value in (None, ""):
                continue
            if len(result) >= 5000:
                raise DeliveryImportError("В одной поставке может быть не более 5000 позиций.")
            if not source_name:
                raise DeliveryImportError(f"Строка {row_number}: не указано исходное наименование.")

            def cell(key):
                index = columns.get(key)
                return row[index] if index is not None and index < len(row) else None

            parameters = _parse_parameters(cell("parameters"))
            part_name = (
                _meaningful(cell("manufacturer_part_name"))
                or _meaningful(cell("english_name"))
                or source_name
            )
            component_type = _meaningful(cell("component_type")) or "Прочее"
            component_type = TYPE_ALIASES.get(component_type, component_type)
            description = source_name if source_name != part_name else ""
            item = {
                "sourceRow": row_number,
                "group": _limited(component_type, row_number, "Тип элемента"),
                "name": _limited(part_name, row_number, "Наименование"),
                "value": _limited(parameters["value"], row_number, "Значение"),
                "unit": _limited(parameters["unit"], row_number, "Единицы измерения"),
                "tol": _limited(parameters["tol"], row_number, "Точность"),
                "description": _limited(description, row_number, "Описание"),
                "case": _limited(parameters["case"], row_number, "Корпус"),
                "manufacturer": _limited(cell("manufacturer"), row_number, "Производитель"),
                "cnt": _parse_quantity(quantity_value, row_number),
                "cellnum": "",
            }
            result.append(item)

        if not result:
            raise DeliveryImportError("В отчёте нет позиций для поставки.")
        return result
    finally:
        workbook.close()
