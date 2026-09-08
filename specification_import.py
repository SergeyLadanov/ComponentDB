"""Import a device specification from a regular or PCB BOM Parser workbook."""
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from delivery_import import (
    _limited, _meaningful, _number_text, _parse_parameters, _parse_quantity,
    _text, _translate_unit, TYPE_ALIASES,
)


class SpecificationImportError(ValueError):
    pass


HEADER_ALIASES = {
    "id компонента": "component_id",
    "id позиции": "component_id",
    "component id": "component_id",
    "классификация": "group",
    "тип": "group",
    "тип элемента": "group",
    "наименование": "name",
    "артикул": "name",
    "наимен. произв.": "name",
    "manufacturer part number": "name",
    "исходное наименование": "source_name",
    "значение": "value",
    "ед. изм.": "unit",
    "единица": "unit",
    "точность": "tol",
    "описание": "description",
    "корпус": "case",
    "производитель": "manufacturer",
    "параметры": "parameters",
    "количество": "quantity",
    "количество на устройство": "quantity",
    "кол-во на устройство": "quantity",
    "qty": "quantity",
}

UNIT_ALIASES = {
    "pf": "пФ", "пф": "пФ",
    "nf": "нФ", "нф": "нФ",
    "uf": "мкФ", "µf": "мкФ", "μf": "мкФ", "мкф": "мкФ",
    "mf": "мФ", "мф": "мФ", "f": "Ф", "ф": "Ф",
    "ohm": "Ом", "ω": "Ом", "ом": "Ом",
    "kohm": "кОм", "kω": "кОм", "ком": "кОм",
    "mohm": "МОм", "mω": "МОм", "мом": "МОм",
    "nh": "нГн", "нгн": "нГн", "uh": "мкГн", "µh": "мкГн",
    "μh": "мкГн", "мкгн": "мкГн", "mh": "мГн", "мгн": "мГн",
    "h": "Гн", "гн": "Гн",
}


def _canonical_unit(value):
    text = _meaningful(value).replace(" ", "")
    return UNIT_ALIASES.get(text.casefold(), _translate_unit(text))


def _compact_parameters(source_name):
    """Recognize passives in compact BOM strings such as 12pF or 10k."""
    text = " ".join(_text(source_name).split())
    result = {"group": "", "value": "", "unit": "", "tol": "", "case": ""}
    tolerance = re.search(r"(?<![\d.,])(\d+(?:[.,]\d+)?)\s*%", text)
    if tolerance:
        result["tol"] = f"{_number_text(tolerance.group(1))}%"
    case_matches = re.findall(r"(?<!\w)(0201|0402|0603|0805|1206|1210|1812)(?!\w)", text, re.IGNORECASE)
    case_letter = re.search(r"\bCase\s+([A-E])\b", text, re.IGNORECASE)
    if case_matches:
        result["case"] = case_matches[-1].upper()
    elif case_letter:
        result["case"] = case_letter.group(1).upper()

    capacitor = re.match(
        r"^(\d+(?:[.,]\d+)?)\s*(pF|nF|uF|µF|μF|mF|F|пФ|нФ|мкФ|мФ|Ф)\b",
        text, re.IGNORECASE,
    )
    if capacitor:
        result.update(group="Конденсатор", value=_number_text(capacitor.group(1)),
                      unit=_canonical_unit(capacitor.group(2)))
        return result

    inductance = re.match(
        r"^(\d+(?:[.,]\d+)?)\s*(nH|uH|µH|μH|mH|H|нГн|мкГн|мГн|Гн)\b",
        text, re.IGNORECASE,
    )
    if inductance:
        result.update(group="Катушка индуктивности", value=_number_text(inductance.group(1)),
                      unit=_canonical_unit(inductance.group(2)))
        return result

    # A bare resistance value is accepted only when the rest of the string looks
    # like a resistor: tolerance plus power rating prevent false positives.
    if tolerance and re.search(r"\b\d+(?:[.,]\d+)?\s*(?:m?W|Вт|мВт)\b", text, re.IGNORECASE):
        resistance = re.match(
            r"^(\d+(?:[.,]\d+)?)\s*(kOhm|MOhm|Ohm|кОм|МОм|Ом|[kKmM])?\b",
            text,
        )
        if resistance:
            suffix = resistance.group(2) or "Ом"
            suffix_units = {"k": "кОм", "K": "кОм", "M": "МОм", "m": "мОм"}
            result.update(group="Резистор", value=_number_text(resistance.group(1)),
                          unit=suffix_units.get(suffix, _canonical_unit(suffix)))
    return result


def parse_specification_workbook(file_stream, filename):
    if Path(filename or "").suffix.lower() != ".xlsx":
        raise SpecificationImportError("Выберите спецификацию в формате .xlsx.")
    try:
        workbook = load_workbook(file_stream, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError):
        raise SpecificationImportError("Не удалось прочитать файл Excel.") from None

    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise SpecificationImportError("В файле нет строки заголовков.")
        columns = {}
        for index, value in enumerate(header):
            key = HEADER_ALIASES.get(_text(value).lower())
            if key and key not in columns:
                columns[key] = index
        # Compact three-column BOMs are also accepted positionally. Some older
        # generators wrote mojibake headers, while the row data stayed intact.
        if len(header) >= 3 and _text(header[0]) == "#":
            columns.setdefault("source_name", 1)
            columns.setdefault("quantity", 2)
        if "quantity" not in columns:
            raise SpecificationImportError(
                "В файле должен быть столбец «Количество на устройство» или «Количество»."
            )
        if not any(key in columns for key in ("component_id", "name", "source_name", "parameters")):
            raise SpecificationImportError(
                "Добавьте ID компонента, наименование или параметры позиции."
            )

        result = []
        for row_number, row in enumerate(rows, start=2):
            def cell(key):
                index = columns.get(key)
                return row[index] if index is not None and index < len(row) else None

            if all(cell(key) in (None, "") for key in columns):
                continue
            if len(result) >= 5000:
                raise SpecificationImportError("В одной спецификации может быть не более 5000 позиций.")
            try:
                quantity = _parse_quantity(cell("quantity"), row_number)
            except ValueError as error:
                raise SpecificationImportError(str(error)) from None

            component_id = _text(cell("component_id"))
            if component_id and (not component_id.isascii() or not component_id.isdigit() or int(component_id) <= 0):
                raise SpecificationImportError(f"Строка {row_number}: некорректный ID компонента.")
            parameters = _parse_parameters(cell("parameters"))
            source_name = _meaningful(cell("source_name"))
            compact = _compact_parameters(source_name)
            name = _meaningful(cell("name")) or source_name
            source_group = _meaningful(cell("group")) or compact["group"] or "Прочее"
            group = TYPE_ALIASES.get(source_group, source_group)
            value = _meaningful(cell("value")) or parameters["value"] or compact["value"]
            unit = _canonical_unit(cell("unit")) if _meaningful(cell("unit")) else (parameters["unit"] or compact["unit"])
            tol = _meaningful(cell("tol")) or parameters["tol"] or compact["tol"]
            case = _meaningful(cell("case")) or parameters["case"] or compact["case"]
            if not component_id and not any((name, value, case)):
                raise SpecificationImportError(f"Строка {row_number}: недостаточно данных для сопоставления позиции.")
            item = {
                "sourceRow": row_number,
                "componentId": int(component_id) if component_id else None,
                "group": _limited(group, row_number, "Классификация"),
                "name": _limited(name, row_number, "Наименование"),
                "value": _limited(_number_text(value) if value else "", row_number, "Значение"),
                "unit": _limited(unit, row_number, "Единицы измерения"),
                "tol": _limited(tol, row_number, "Точность"),
                "description": _limited(_meaningful(cell("description")) or (source_name if compact["group"] else ""), row_number, "Описание"),
                "case": _limited(case, row_number, "Корпус"),
                "manufacturer": _limited(_meaningful(cell("manufacturer")), row_number, "Производитель"),
                "quantityPerDevice": quantity,
            }
            result.append(item)
        if not result:
            raise SpecificationImportError("В файле нет позиций спецификации.")
        return result
    finally:
        workbook.close()
