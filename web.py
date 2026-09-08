#!/usr/bin/env python3
from flask import Flask, render_template, request, redirect, session, url_for, send_file
from datetime import timedelta
from functools import wraps
from io import BytesIO
import os
from pathlib import Path
import secrets
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.exceptions import RequestEntityTooLarge
from delivery_import import DeliveryImportError, parse_delivery_workbook
from specification_import import SpecificationImportError, parse_specification_workbook
from version import __version__


# Текущий путь приложения
path = os.path.dirname(os.path.abspath(__file__))
try:
    from config import HTTP_HOST
    from config import HTTP_PORT
    from config import ACCOUNTS
except ModuleNotFoundError as error:
    if error.name != 'config':
        raise
    print("This is the first start of application")
    config_content =  '''#----WEB server settings----#
# WEB server host
HTTP_HOST = "localhost"
# WEB server port
HTTP_PORT = 5000
# Userts and passwords
ACCOUNTS = ["user1:pswd1", "user2:pswd2"]
#----Database settings----#
# Host for mySQL database
DB_HOST = "localhost"
# Port for mySQL database
DB_PORT = 3306
# User of database
DB_USER = ""
# Name of database
DB_NAME = ""
# Password for database
DB_PSWD = ""
#-----Backup settings-----#
# Using relative path
RELATIVE_PATH = True
# Path for saving dump
DUMP_PATH = "/dump/dump.sql"'''

    with open(os.path.join(path, 'config.py'), 'w', encoding='utf-8') as config_file:
        config_file.write(config_content)
    print("Config file was created, please config application and restart it")
    exit(0)

# Импорт модуля для работы с БД
import db_if

# Учетные записи общие для страницы входа и API.
def checkAuth(username, password):
    if not isinstance(username, str) or not isinstance(password, str):
        return False
    for item in ACCOUNTS:
        usr, separator, pswd = item.partition(':')
        if separator and usr == username and secrets.compare_digest(pswd.encode(), password.encode()):
            return True
    return False


def session_username():
    username = session.get('username')
    return username if username and any(item.partition(':')[0] == username for item in ACCOUNTS) else None


def valid_csrf():
    token = session.get('csrf_token')
    supplied = request.headers.get('X-CSRF-Token', '')
    return bool(token and secrets.compare_digest(token.encode(), supplied.encode()))


# Декоратор авторизации
def auth_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if session_username():
            if request.method not in ('GET', 'HEAD', 'OPTIONS') and not valid_csrf():
                return {'error': 'Сессия изменилась. Обновите страницу и повторите попытку.'}, 403
            return f(*args, **kwargs)
        # Basic Auth остается доступен для существующих клиентов API.
        auth = request.authorization
        if request.endpoint != 'index' and auth and auth.type == 'basic' and checkAuth(auth.username, auth.password):
            return f(*args, **kwargs)
        if request.endpoint == 'index':
            return redirect(url_for('login'))
        return {'error': 'Требуется авторизация.'}, 401
    return decorated
#------------------------------
app = Flask(__name__)
# Учитываем префикс, переданный доверенным reverse proxy в
# X-Forwarded-Prefix (например, /components).
app.wsgi_app = ProxyFix(
    app.wsgi_app,
    x_for=1,
    x_proto=1,
    x_host=1,
    x_prefix=1,
)
app.config.update(
    SECRET_KEY=os.environ.get('COMPONENTDB_SECRET_KEY') or secrets.token_hex(32),
    SESSION_COOKIE_NAME='componentdb_session',
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.environ.get('COMPONENTDB_COOKIE_SECURE') == '1',
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
    SESSION_REFRESH_EACH_REQUEST=False,
    MAX_CONTENT_LENGTH=5 * 1024 * 1024,
)


@app.after_request
def prevent_private_caching(response):
    if request.endpoint != 'static':
        response.headers['Cache-Control'] = 'no-store'
    return response


@app.errorhandler(RequestEntityTooLarge)
def upload_too_large(_error):
    return {'error': 'Размер Excel-файла не должен превышать 5 МБ.'}, 413


@app.route('/auth/session')
def auth_session():
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_urlsafe(32)
    return {
        'username': session_username(),
        'csrfToken': session['csrf_token'],
        'version': __version__,
    }


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'GET':
        if session_username():
            return redirect(url_for('index'))
        return render_template('index.html')
    if not valid_csrf():
        return {'error': 'Сессия изменилась. Обновите страницу и повторите попытку.'}, 403
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or not checkAuth(data.get('username'), data.get('password')):
        return {'error': 'Неверный логин или пароль.'}, 401
    session.clear()
    session.permanent = True
    session['username'] = data['username']
    session['csrf_token'] = secrets.token_urlsafe(32)
    return {'username': session['username']}


@app.route('/logout', methods=['POST'])
def logout():
    if not valid_csrf():
        return {'error': 'Сессия изменилась. Обновите страницу и повторите попытку.'}, 403
    session.clear()
    return {'ok': True}

# Корневой каталог
@app.route("/")
@auth_required
def index():
    return render_template('index.html')


# Чтение базы данных
@app.route('/get_data', methods=['GET', 'POST'])
@auth_required
def control():
    type_data = request.args.get('filter', '')
    data = db_if.getData(type_data)
    return data


def positive_id(value, label):
    value = str(value or '')
    if not value.isascii() or not value.isdigit() or len(value) > 19 or int(value) <= 0:
        raise ValueError(f'Некорректный ID {label}.')
    return int(value)


def validate_component_data(data, positive_quantity=False):
    if not isinstance(data, dict):
        raise ValueError('Некорректные данные позиции.')
    result = {
        key: data.get(key)
        for key in ('group', 'name', 'value', 'unit', 'tol', 'description',
                    'case', 'manufacturer', 'cnt', 'cellnum')
    }
    quantity = str(result['cnt'] if result['cnt'] is not None else '')
    if (not quantity.isascii() or not quantity.isdigit() or len(quantity) > 16
            or int(quantity) > 9007199254740991
            or (positive_quantity and int(quantity) <= 0)):
        qualifier = 'положительным ' if positive_quantity else 'неотрицательным '
        raise ValueError(f'Количество должно быть целым {qualifier}числом.')
    if not result['group']:
        raise ValueError('Укажите тип компонента.')
    result['cnt'] = int(quantity)
    for key in ('group', 'name', 'value', 'unit', 'tol', 'description',
                'case', 'manufacturer', 'cellnum'):
        result[key] = str(result[key] or '').strip()
        if len(result[key]) > 255:
            raise ValueError('Значение поля не должно превышать 255 символов.')
    return result


def validate_specification_header(data):
    if not isinstance(data, dict):
        raise ValueError('Некорректные данные спецификации.')
    name = str(data.get('name') or '').strip()
    if not name:
        raise ValueError('Укажите название спецификации.')
    if len(name) > 255:
        raise ValueError('Название спецификации длиннее 255 символов.')
    device_quantity = str(data.get('deviceQuantity') or '')
    if (not device_quantity.isascii() or not device_quantity.isdigit()
            or len(device_quantity) > 16 or int(device_quantity) <= 0
            or int(device_quantity) > 9007199254740991):
        raise ValueError('Количество устройств должно быть целым положительным числом.')
    return name, int(device_quantity)


def validate_specification_item(data):
    if not isinstance(data, dict):
        raise ValueError('Некорректные данные позиции спецификации.')
    quantity = str(data.get('quantityPerDevice') or '')
    if (not quantity.isascii() or not quantity.isdigit() or len(quantity) > 16
            or int(quantity) <= 0 or int(quantity) > 9007199254740991):
        raise ValueError('Количество на устройство должно быть целым положительным числом.')
    component_id = data.get('componentId')
    component_id = positive_id(component_id, 'компонента') if str(component_id or '').strip() else None
    result = {'componentId': component_id, 'quantityPerDevice': int(quantity)}
    for key in ('group', 'name', 'value', 'unit', 'tol', 'description', 'case', 'manufacturer'):
        result[key] = str(data.get(key) or '').strip()
        if len(result[key]) > 255:
            raise ValueError('Значение поля не должно превышать 255 символов.')
    if component_id is None and not result['group']:
        raise ValueError('Укажите классификацию или ID складского компонента.')
    return result


@app.route('/specifications', methods=['GET', 'POST'])
@auth_required
def specifications():
    if request.method == 'GET':
        return {'data': db_if.getSpecifications()}
    try:
        name, quantity = validate_specification_header(request.get_json(silent=True))
    except ValueError as error:
        return {'error': str(error)}, 400
    return {'id': db_if.createSpecification(name, quantity)}, 201


@app.route('/specifications/import', methods=['POST'])
@auth_required
def import_specification():
    upload = request.files.get('file')
    if upload is None or not upload.filename:
        return {'error': 'Выберите Excel-файл спецификации.'}, 400
    source_file = upload.filename.replace('\\', '/').rsplit('/', 1)[-1]
    try:
        name, quantity = validate_specification_header({
            'name': request.form.get('name') or Path(source_file).stem,
            'deviceQuantity': request.form.get('deviceQuantity') or '1',
        })
        if len(source_file) > 255:
            raise ValueError('Имя файла длиннее 255 символов.')
        items = parse_specification_workbook(upload.stream, source_file)
    except (ValueError, SpecificationImportError) as error:
        return {'error': str(error)}, 400
    specification_id = db_if.createSpecification(name, quantity, source_file, items)
    return {'id': specification_id, 'items': len(items)}, 201


@app.route('/specifications/<specification_id>', methods=['PUT', 'DELETE'])
@auth_required
def specification_detail(specification_id):
    try:
        specification_id = positive_id(specification_id, 'спецификации')
        if request.method == 'DELETE':
            if not db_if.deleteSpecification(specification_id):
                return {'error': 'Спецификация не найдена.'}, 404
            return {'ok': True}
        name, quantity = validate_specification_header(request.get_json(silent=True))
    except ValueError as error:
        return {'error': str(error)}, 400
    if not db_if.updateSpecification(specification_id, name, quantity):
        return {'error': 'Спецификация не найдена.'}, 404
    return {'ok': True}


@app.route('/specifications/<specification_id>/items', methods=['POST'])
@auth_required
def specification_items(specification_id):
    try:
        specification_id = positive_id(specification_id, 'спецификации')
        data = validate_specification_item(request.get_json(silent=True))
    except ValueError as error:
        return {'error': str(error)}, 400
    item_id = db_if.addSpecificationItem(specification_id, data)
    if item_id is None:
        return {'error': 'Спецификация не найдена.'}, 404
    return {'id': item_id}, 201


@app.route('/specifications/<specification_id>/items/<item_id>', methods=['PUT', 'DELETE'])
@auth_required
def specification_item(specification_id, item_id):
    try:
        specification_id = positive_id(specification_id, 'спецификации')
        item_id = positive_id(item_id, 'позиции')
        if request.method == 'DELETE':
            if not db_if.deleteSpecificationItem(specification_id, item_id):
                return {'error': 'Позиция спецификации не найдена.'}, 404
            return {'ok': True}
        data = validate_specification_item(request.get_json(silent=True))
    except ValueError as error:
        return {'error': str(error)}, 400
    if not db_if.updateSpecificationItem(specification_id, item_id, data):
        return {'error': 'Позиция спецификации не найдена.'}, 404
    return {'ok': True}


@app.route('/specifications/<specification_id>/export', methods=['GET'])
@auth_required
def export_specification(specification_id):
    try:
        specification_id = positive_id(specification_id, 'спецификации')
    except ValueError as error:
        return {'error': str(error)}, 400
    include_all = request.args.get('scope') == 'all'
    export_data = db_if.getSpecificationForExport(specification_id, include_all)
    if export_data is None:
        return {'error': 'Спецификация не найдена.'}, 404
    specification, rows = export_data
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Дозаказ' if not include_all else 'Потребность'
    headers = ['#', 'Тип элемента', 'Наименование', 'Описание', 'Количество']
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='18746C')
    for number, row in enumerate(rows, start=1):
        name_parts = []
        for value in (
            row['name'], row['value'], row['unit'], row['tol'],
            row['case'], row['manufacturer'],
        ):
            value = ' '.join(str(value or '').split())
            if value and value != '-':
                name_parts.append(value)
        sheet.append([
            number,
            row['group'],
            ' '.join(name_parts),
            row['description'],
            row['required'] if include_all else row['toOrder'],
        ])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    widths = [8, 24, 60, 50, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = 'вся-потребность' if include_all else 'дозаказ'
    filename = f'{specification.Name}-{suffix}.xlsx'
    response = send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='specification.xlsx',
    )
    response.headers['Content-Disposition'] = (
        f"attachment; filename=specification.xlsx; filename*=UTF-8''{quote(filename)}"
    )
    return response


@app.route('/specifications/<specification_id>/delivery', methods=['POST'])
@auth_required
def create_delivery_from_specification(specification_id):
    try:
        specification_id = positive_id(specification_id, 'спецификации')
    except ValueError as error:
        return {'error': str(error)}, 400
    result = db_if.createExpectedDeliveryFromSpecification(specification_id)
    if result is None:
        return {'error': 'Спецификация не найдена.'}, 404
    if not result['items']:
        return {'error': 'Дозаказ больше не требуется. Обновите спецификацию.'}, 409
    return result, 201


@app.route('/deliveries', methods=['GET'])
@auth_required
def deliveries():
    return {'data': db_if.getExpectedDeliveries()}


@app.route('/deliveries/import', methods=['POST'])
@auth_required
def import_delivery():
    upload = request.files.get('file')
    if upload is None or not upload.filename:
        return {'error': 'Выберите Excel-отчёт.'}, 400
    source_file = upload.filename.replace('\\', '/').rsplit('/', 1)[-1]
    name = str(request.form.get('name') or Path(source_file).stem).strip()
    if not name:
        name = 'Поставка'
    if len(name) > 255 or len(source_file) > 255:
        return {'error': 'Название поставки или файла длиннее 255 символов.'}, 400
    try:
        items = parse_delivery_workbook(upload.stream, source_file)
    except DeliveryImportError as error:
        return {'error': str(error)}, 400
    delivery_id = db_if.createExpectedDelivery(name, source_file, items)
    return {'id': delivery_id, 'items': len(items)}, 201


@app.route('/deliveries/<delivery_id>/items', methods=['PUT'])
@auth_required
def update_delivery_items(delivery_id):
    try:
        delivery_id = positive_id(delivery_id, 'поставки')
        payload = request.get_json(silent=True)
        source_items = payload.get('items') if isinstance(payload, dict) else None
        if not isinstance(source_items, list) or not source_items or len(source_items) > 5000:
            raise ValueError('Передайте от 1 до 5000 позиций для сохранения.')
        items = []
        item_ids = set()
        for source_item in source_items:
            item_id = positive_id(
                source_item.get('id') if isinstance(source_item, dict) else None,
                'позиции',
            )
            if item_id in item_ids:
                raise ValueError('Список содержит повторяющийся ID позиции.')
            item_ids.add(item_id)
            item = validate_component_data(source_item, positive_quantity=True)
            item['id'] = item_id
            items.append(item)
        saved = db_if.updateExpectedDeliveryItems(delivery_id, items)
    except ValueError as error:
        return {'error': str(error)}, 400
    if not saved:
        return {'error': 'Поставка не найдена.'}, 404
    return {'updated': len(items)}


@app.route('/deliveries/<delivery_id>/items/<item_id>', methods=['PUT', 'DELETE'])
@auth_required
def update_delivery_item(delivery_id, item_id):
    try:
        delivery_id = positive_id(delivery_id, 'поставки')
        item_id = positive_id(item_id, 'позиции')
        if request.method == 'DELETE':
            if not db_if.deleteExpectedDeliveryItem(delivery_id, item_id):
                return {'error': 'Поставка или позиция не найдена.'}, 404
            return {'ok': True}
        data = validate_component_data(request.get_json(silent=True), positive_quantity=True)
    except ValueError as error:
        return {'error': str(error)}, 400
    if not db_if.updateExpectedDeliveryItem(delivery_id, item_id, data):
        return {'error': 'Поставка или позиция не найдена.'}, 404
    return {'ok': True}


@app.route('/deliveries/<delivery_id>/confirm', methods=['POST'])
@auth_required
def confirm_delivery(delivery_id):
    try:
        delivery_id = positive_id(delivery_id, 'поставки')
        result = db_if.confirmExpectedDelivery(delivery_id)
    except ValueError as error:
        return {'error': str(error)}, 400
    if result is None:
        return {'error': 'Поставка уже обработана или не найдена.'}, 409
    return result


@app.route('/deliveries/<delivery_id>/cancel', methods=['POST'])
@auth_required
def cancel_delivery(delivery_id):
    try:
        delivery_id = positive_id(delivery_id, 'поставки')
    except ValueError as error:
        return {'error': str(error)}, 400
    if not db_if.cancelExpectedDelivery(delivery_id):
        return {'error': 'Поставка уже обработана или не найдена.'}, 409
    return {'ok': True}

# Обработка запроса на изменение базы данных
@app.route('/request_handler', methods=['GET', 'POST'])
@auth_required
def request_handler():
    type_request = request.form.get('reqtype')
    if type_request not in ('Add', 'Edit', 'Remove'):
        return {'error': 'Неизвестная операция.'}, 400
    repply = "OK"
    row_data = {
        "id": request.form.get('id'), 
        "group": request.form.get('group'),
        "name": request.form.get('name'),
        "value": request.form.get('value'),
        "unit": request.form.get('unit'),
        "tol": request.form.get('tol'),
        "description": request.form.get('description'),
        "case": request.form.get('case'),
        "manufacturer": request.form.get('manufacturer'),
        "cnt": request.form.get('cnt'),
        "cellnum": request.form.get('cellnum')
    }

    if type_request in ('Edit', 'Remove'):
        row_id = row_data['id'] or ''
        if not row_id.isascii() or not row_id.isdigit() or len(row_id) > 19 or int(row_id) <= 0:
            return {'error': 'Некорректный ID позиции.'}, 400

    if type_request != 'Remove':
        try:
            row_data.update(validate_component_data(row_data))
        except ValueError as error:
            return {'error': str(error)}, 400

    if type_request == "Add":
        repply = db_if.addPosition(row_data)
    if type_request == "Remove":
        repply = db_if.removePosition(row_data)
    if type_request == "Edit":
        repply = db_if.editPosition(row_data)

    return repply
    


# Запуск приложения
if __name__ == '__main__':
    app.run(host=HTTP_HOST, port=HTTP_PORT)
