import base64
from datetime import timedelta
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from tests.support import component, load_test_app


class ComponentApiTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.app, self.db = load_test_app(Path(self.directory.name) / 'components.sqlite')
        self.client = self.app.test_client()
        token = base64.b64encode(b'tester:testing').decode()
        self.headers = {'Authorization': f'Basic {token}'}

    def tearDown(self):
        self.db.dbhandle.close()
        self.directory.cleanup()

    def post(self, operation, **overrides):
        return self.client.post('/request_handler', data=component(reqtype=operation, **overrides), headers=self.headers)

    def rows(self, query=''):
        response = self.client.get('/get_data' + query, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        return response.json['data']

    def login(self, username='tester', password='testing'):
        token = self.client.get('/auth/session').json['csrfToken']
        return self.client.post('/login', json={'username': username, 'password': password},
                                headers={'X-CSRF-Token': token})

    def session_headers(self):
        return {'X-CSRF-Token': self.client.get('/auth/session').json['csrfToken']}

    def delivery_report(self, rows=None, headers=None):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers or [
            '#', 'Исходное наименование', 'Тип элемента', 'Параметры',
            'Наимен. произв.', 'Производитель', 'Количество',
        ])
        for row in rows or [[
            1, '10 кОм 1% 0603', 'Резистор',
            'Способ монтажа: SMD\nЗначение: 10.0 кОм\nКорпус: 0603\nТочность: 1.0 %\nТип: -',
            'RC0603', 'Yageo', 25,
        ]]:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def import_delivery(self, report=None, name='Заказ 42'):
        return self.client.post(
            '/deliveries/import',
            data={'name': name, 'file': (report or self.delivery_report(), 'ResultTable.xlsx')},
            content_type='multipart/form-data',
            headers=self.headers,
        )

    def specification_report(self, rows=None):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            'Тип элемента', 'Параметры', 'Наимен. произв.',
            'Производитель', 'Количество на устройство',
        ])
        for row in rows or [[
            'Резистор', 'Значение: 10.0 кОм\nКорпус: 0603\nТочность: 1.0 %',
            '-', '', 2,
        ]]:
            worksheet.append(row)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def compact_specification_report(self, rows=None):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(['#', '�������� ������������', '����������'])
        for index, row in enumerate(rows or [
            ['100nF 5% 16V X7R 0603', 2],
            ['10k 5% 0.063W 0603', 8],
            ['STM32F103C8T6', 1],
        ], start=1):
            worksheet.append([index, *row])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def test_auth_required_for_page_read_and_write(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, '/login')
        for path in ('/get_data', '/request_handler'):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 401)
            self.assertNotIn('WWW-Authenticate', response.headers)
        response = self.client.post('/request_handler', data=component(reqtype='Add'))
        self.assertEqual(response.status_code, 401)
        self.assertEqual(self.rows(), [])

    def test_page_mounts_react_without_legacy_scripts(self):
        self.login()
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('id="app"', html)
        self.assertIn('/static/dist/index.js', html)
        self.assertNotIn('jquery', html)

    def test_reverse_proxy_prefix_is_used_for_redirects_and_assets(self):
        headers = {'X-Forwarded-Prefix': '/components'}
        response = self.client.get('/', headers=headers)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.location, '/components/login')

        html = self.client.get('/login', headers=headers).get_data(as_text=True)
        self.assertIn('content="/components"', html)
        self.assertIn('/components/static/favicon.ico', html)
        self.assertIn('/components/static/dist/index.js', html)

    def test_login_page_and_session_lifecycle(self):
        self.assertEqual(self.client.get('/login').status_code, 200)
        anonymous = self.client.get('/auth/session')
        self.assertIsNone(anonymous.json['username'])
        self.assertEqual(anonymous.headers['Cache-Control'], 'no-store')
        self.assertEqual(self.login().status_code, 200)
        session_info = self.client.get('/auth/session').json
        self.assertEqual(session_info['username'], 'tester')
        self.assertNotEqual(session_info['csrfToken'], anonymous.json['csrfToken'])
        self.assertEqual(self.client.get('/login').location, '/')
        self.assertEqual(self.client.get('/').status_code, 200)
        self.assertEqual(self.client.get('/get_data').status_code, 200)
        response = self.client.post('/request_handler', data=component(reqtype='Add'), headers=self.session_headers())
        self.assertEqual(response.text, 'True')
        self.assertEqual(len(self.client.get('/get_data').json['data']), 1)
        self.assertEqual(self.client.post('/logout', headers=self.session_headers()).status_code, 200)
        self.assertEqual(self.client.get('/').location, '/login')
        self.assertEqual(self.client.get('/get_data').status_code, 401)
        self.assertEqual(self.client.post('/request_handler', data=component(reqtype='Remove', id='1')).status_code, 401)
        self.assertIsNone(self.client.get('/auth/session').json['username'])

    def test_invalid_credentials_do_not_create_session(self):
        for username, password in [('tester', 'wrong'), ('unknown', 'testing'), ('', ''), (None, []), ('tester', 'неверный')]:
            with self.subTest(username=username, password=password):
                response = self.login(username, password)
                self.assertEqual(response.status_code, 401)
                self.assertEqual(response.json['error'], 'Неверный логин или пароль.')
                self.assertNotIn('WWW-Authenticate', response.headers)
                self.assertIsNone(self.client.get('/auth/session').json['username'])
        self.assertEqual(self.client.post('/login', json=['tester', 'testing'], headers=self.session_headers()).status_code, 401)

    def test_csrf_required_for_login_logout_and_session_writes(self):
        self.client.get('/auth/session')
        self.assertEqual(self.client.post('/login', json={'username': 'tester', 'password': 'testing'}).status_code, 403)
        self.login()
        for headers in ({}, {'X-CSRF-Token': 'incorrect'}):
            self.assertEqual(self.client.post('/logout', headers=headers).status_code, 403)
            self.assertEqual(self.client.post('/request_handler', data=component(reqtype='Add'), headers=headers).status_code, 403)
            self.assertEqual(self.client.post('/deliveries/import', headers=headers).status_code, 403)
        self.assertEqual(self.client.get('/logout').status_code, 405)
        self.assertEqual(self.client.get('/get_data').json['data'], [])
        self.assertEqual(self.client.get('/auth/session').json['username'], 'tester')

    def test_cookie_flags_and_session_expiration(self):
        response = self.login()
        cookie = response.headers['Set-Cookie']
        self.assertIn('HttpOnly', cookie)
        self.assertIn('SameSite=Lax', cookie)
        self.assertNotIn('testing', cookie)
        self.app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(seconds=-1)
        self.assertEqual(self.client.get('/get_data').status_code, 401)
        self.assertEqual(self.client.get('/').location, '/login')

    def test_tampered_cookie_cannot_authenticate(self):
        self.login()
        self.client.set_cookie(self.app.config['SESSION_COOKIE_NAME'], 'forged-session')
        self.assertEqual(self.client.get('/get_data').status_code, 401)
        self.assertEqual(self.client.get('/').location, '/login')

    def test_add_filter_edit_write_off_and_remove(self):
        self.assertEqual(self.post('Add').text, 'True')
        row = self.rows()[0]
        self.assertEqual(len(row), 12)
        self.assertEqual(row[1:11], ['Резистор', 'RC0603', '10', 'кОм', '1%', 'Тестовый компонент', '0603', 'Yageo', 5, 'A-01'])
        self.assertEqual(len(self.rows('?filter=Резистор')), 1)
        self.assertEqual(self.rows('?filter=Конденсатор'), [])
        response = self.post('Edit', id=str(row[0]), cnt='0', cellnum='B-02')
        self.assertRegex(response.text, r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$')
        self.assertEqual(self.rows()[0][9:11], [0, 'B-02'])
        self.assertEqual(self.post('Remove', id=str(row[0])).text, 'True')
        self.assertEqual(self.rows(), [])

    def test_add_duplicate_merges_quantities(self):
        self.post('Add')
        self.assertEqual(self.post('Add', cnt='2').text, 'Match')
        self.assertEqual(len(self.rows()), 1)
        self.assertEqual(self.rows()[0][9], 7)

    def test_edit_duplicate_merges_rows(self):
        self.post('Add')
        self.post('Add', name='Another part', cnt='3')
        row_id = self.rows()[1][0]
        self.assertEqual(self.post('Edit', id=str(row_id), cnt='3').text, 'Match')
        self.assertEqual(len(self.rows()), 1)
        self.assertEqual(self.rows()[0][9], 8)

    def test_invalid_quantities_cannot_modify_database(self):
        for value in ('-1', '1.5', '', 'abc', '9007199254740992', '1' * 100, '²'):
            with self.subTest(value=value):
                response = self.post('Add', cnt=value)
                self.assertEqual(response.status_code, 400)
                self.assertIn('error', response.json)
        self.assertEqual(self.rows(), [])

    def test_invalid_operation_id_and_long_fields(self):
        self.assertEqual(self.post('Unknown').status_code, 400)
        self.assertEqual(self.post('Edit', id='invalid').status_code, 400)
        self.assertEqual(self.post('Remove', id='0').status_code, 400)
        self.assertEqual(self.post('Add', name='a' * 256).status_code, 400)
        self.assertEqual(self.post('Add', group='').status_code, 400)
        self.assertEqual(self.rows(), [])

    def test_missing_position_reports_failure(self):
        self.assertEqual(self.post('Edit', id='999').text, 'False')
        self.assertEqual(self.post('Remove', id='999').text, 'False')

    def test_excel_report_creates_editable_delivery_and_confirm_adds_stock(self):
        response = self.import_delivery()
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json['items'], 1)

        delivery = self.client.get('/deliveries', headers=self.headers).json['data'][0]
        item = delivery['items'][0]
        self.assertEqual(delivery['name'], 'Заказ 42')
        self.assertEqual(item['sourceRow'], 2)
        self.assertEqual(
            [item['group'], item['name'], item['value'], item['unit'], item['tol'], item['case'], item['manufacturer'], item['cnt'], item['cellnum']],
            ['Резистор', 'RC0603', '10', 'кОм', '1%', '0603', 'Yageo', '25', ''],
        )

        item['cellnum'] = 'B-14'
        item['description'] = 'Партия для стенда'
        updated = self.client.put(
            f"/deliveries/{delivery['id']}/items/{item['id']}",
            json={key: item[key] for key in ('group', 'name', 'value', 'unit', 'tol', 'description', 'case', 'manufacturer', 'cnt', 'cellnum')},
            headers=self.headers,
        )
        self.assertEqual(updated.status_code, 200)

        confirmed = self.client.post(
            f"/deliveries/{delivery['id']}/confirm", headers=self.headers
        )
        self.assertEqual(confirmed.status_code, 200)
        self.assertEqual(confirmed.json, {'created': 1, 'merged': 0, 'items': 1})
        self.assertEqual(self.client.get('/deliveries', headers=self.headers).json['data'], [])
        self.assertEqual(
            self.rows()[0][1:11],
            ['Резистор', 'RC0603', '10', 'кОм', '1%', 'Партия для стенда', '0603', 'Yageo', 25, 'B-14'],
        )
        repeated = self.client.post(
            f"/deliveries/{delivery['id']}/confirm", headers=self.headers
        )
        self.assertEqual(repeated.status_code, 409)
        self.assertEqual(self.rows()[0][9], 25)

    def test_delivery_can_be_confirmed_without_cell_number(self):
        delivery_id = self.import_delivery().json['id']
        confirmed = self.client.post(f'/deliveries/{delivery_id}/confirm', headers=self.headers)
        self.assertEqual(confirmed.status_code, 200)
        self.assertEqual(confirmed.json, {'created': 1, 'merged': 0, 'items': 1})
        self.assertEqual(self.rows()[0][10], '')

    def test_delivery_items_can_be_saved_together_and_deleted(self):
        report = self.delivery_report(rows=[
            [1, '10 кОм 1% 0603', 'Резистор', 'Значение: 10.0 кОм', 'RC0603', 'Yageo', 25],
            [2, '100 нФ 10% 0603', 'Конденсатор', 'Значение: 100.0 nF', 'CC0603', 'Murata', 30],
        ])
        delivery_id = self.import_delivery(report).json['id']
        delivery = self.client.get('/deliveries', headers=self.headers).json['data'][0]
        changed_items = []
        for index, item in enumerate(delivery['items'], start=1):
            item['cellnum'] = f'C-{index:02}'
            item['description'] = f'Строка {index}'
            changed_items.append({
                key: item[key]
                for key in ('id', 'group', 'name', 'value', 'unit', 'tol',
                            'description', 'case', 'manufacturer', 'cnt', 'cellnum')
            })
        saved = self.client.put(
            f'/deliveries/{delivery_id}/items',
            json={'items': changed_items},
            headers=self.headers,
        )
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json, {'updated': 2})
        saved_items = self.client.get('/deliveries', headers=self.headers).json['data'][0]['items']
        self.assertEqual([item['cellnum'] for item in saved_items], ['C-01', 'C-02'])

        deleted = self.client.delete(
            f"/deliveries/{delivery_id}/items/{saved_items[0]['id']}",
            headers=self.headers,
        )
        self.assertEqual(deleted.status_code, 200)
        remaining = self.client.get('/deliveries', headers=self.headers).json['data'][0]['items']
        self.assertEqual(len(remaining), 1)
        self.assertEqual(remaining[0]['cellnum'], 'C-02')

    def test_cancel_delivery_keeps_stock_unchanged(self):
        delivery_id = self.import_delivery().json['id']
        response = self.client.post(f'/deliveries/{delivery_id}/cancel', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get('/deliveries', headers=self.headers).json['data'], [])
        self.assertEqual(self.rows(), [])

    def test_confirm_delivery_merges_matching_stock_position(self):
        self.assertEqual(self.post('Add', description='10 кОм 1% 0603').status_code, 200)
        delivery_id = self.import_delivery().json['id']
        item = self.client.get('/deliveries', headers=self.headers).json['data'][0]['items'][0]
        item['cellnum'] = 'B-14'
        response = self.client.put(
            f"/deliveries/{delivery_id}/items/{item['id']}",
            json={key: item[key] for key in ('group', 'name', 'value', 'unit', 'tol', 'description', 'case', 'manufacturer', 'cnt', 'cellnum')},
            headers=self.headers,
        )
        self.assertEqual(response.status_code, 200)
        confirmed = self.client.post(f'/deliveries/{delivery_id}/confirm', headers=self.headers)
        self.assertEqual(confirmed.json, {'created': 0, 'merged': 1, 'items': 1})
        self.assertEqual(self.rows()[0][9:11], [30, 'A-01'])

    def test_minimal_report_columns_are_supported(self):
        report = self.delivery_report(
            headers=['#', 'Исходное наименование', 'Количество'],
            rows=[[1, 'STM32H743ZIT6', '2']],
        )
        self.assertEqual(self.import_delivery(report).status_code, 201)
        item = self.client.get('/deliveries', headers=self.headers).json['data'][0]['items'][0]
        self.assertEqual(item['group'], 'Прочее')
        self.assertEqual(item['name'], 'STM32H743ZIT6')
        self.assertEqual(item['cnt'], '2')

    def test_invalid_xlsx_is_reported_without_creating_delivery(self):
        response = self.import_delivery(BytesIO(b'not an xlsx file'))
        self.assertEqual(response.status_code, 400)
        self.assertIn('Не удалось прочитать', response.json['error'])
        self.assertEqual(self.client.get('/deliveries', headers=self.headers).json['data'], [])

    def test_english_units_are_translated_during_import(self):
        source_units = ['Ohm', 'kOhm', 'MOhm', 'pF', 'uF', 'nH', 'uH', 'H']
        expected_units = ['Ом', 'кОм', 'МОм', 'пФ', 'мкФ', 'нГн', 'мкГн', 'Гн']
        report = self.delivery_report(
            headers=['#', 'Исходное наименование', 'Тип элемента', 'Параметры', 'Количество'],
            rows=[
                [index, f'Part {index}', 'Прочее', f'Значение: 1.0 {unit}', 1]
                for index, unit in enumerate(source_units, start=1)
            ],
        )
        self.assertEqual(self.import_delivery(report).status_code, 201)
        items = self.client.get('/deliveries', headers=self.headers).json['data'][0]['items']
        self.assertEqual([item['unit'] for item in items], expected_units)

    def test_specification_calculates_shortage_and_reacts_to_device_count(self):
        self.post('Add', cnt='5', description='')
        created = self.client.post(
            '/specifications', json={'name': 'Плата управления', 'deviceQuantity': 3},
            headers=self.headers,
        )
        self.assertEqual(created.status_code, 201)
        specification_id = created.json['id']
        added = self.client.post(
            f'/specifications/{specification_id}/items',
            json={'componentId': '1', 'quantityPerDevice': '2'},
            headers=self.headers,
        )
        self.assertEqual(added.status_code, 201)
        specification = self.client.get('/specifications', headers=self.headers).json['data'][0]
        item = specification['items'][0]
        self.assertEqual(item['requiredQuantity'], '6')
        self.assertEqual(item['stockQuantity'], '5')
        self.assertEqual(item['shortageQuantity'], '1')
        self.assertEqual(item['status'], 'shortage')
        self.assertEqual(specification['summary']['toOrder'], 1)

        changed = self.client.put(
            f'/specifications/{specification_id}',
            json={'name': 'Плата управления', 'deviceQuantity': 2},
            headers=self.headers,
        )
        self.assertEqual(changed.status_code, 200)
        item = self.client.get('/specifications', headers=self.headers).json['data'][0]['items'][0]
        self.assertEqual(item['requiredQuantity'], '4')
        self.assertEqual(item['status'], 'enough')

    def test_specification_import_matches_passive_parameters_and_exports_orders(self):
        self.post('Add', cnt='3', description='')
        response = self.client.post(
            '/specifications/import',
            data={
                'name': 'Импорт BOM', 'deviceQuantity': '2',
                'file': (self.specification_report(), 'bom.xlsx'),
            },
            content_type='multipart/form-data', headers=self.headers,
        )
        self.assertEqual(response.status_code, 201)
        specification_id = response.json['id']
        specification = self.client.get('/specifications', headers=self.headers).json['data'][0]
        item = specification['items'][0]
        self.assertEqual(item['componentId'], '1')
        self.assertEqual(item['requiredQuantity'], '4')
        self.assertEqual(item['shortageQuantity'], '1')

        exported = self.client.get(
            f'/specifications/{specification_id}/export', headers=self.headers,
        )
        self.assertEqual(exported.status_code, 200)
        workbook = __import__('openpyxl').load_workbook(BytesIO(exported.data), data_only=True)
        values = list(workbook.active.values)
        self.assertEqual(values, [
            ('#', 'Тип элемента', 'Наименование', 'Описание', 'Количество'),
            (1, 'Резистор', 'RC0603 10 кОм 1% 0603 Yageo', None, 1),
        ])
        workbook.close()

        exported = self.client.get(
            f'/specifications/{specification_id}/export?scope=all', headers=self.headers,
        )
        workbook = __import__('openpyxl').load_workbook(BytesIO(exported.data), data_only=True)
        values = list(workbook.active.values)
        self.assertEqual(values, [
            ('#', 'Тип элемента', 'Наименование', 'Описание', 'Количество'),
            (1, 'Резистор', 'RC0603 10 кОм 1% 0603 Yageo', None, 4),
        ])
        workbook.close()

    def test_reorder_list_creates_expected_delivery(self):
        self.post('Add', cnt='3', description='')
        created = self.client.post(
            '/specifications', json={'name': 'Плата управления', 'deviceQuantity': 2},
            headers=self.headers,
        )
        specification_id = created.json['id']
        for quantity in ('1', '1'):
            response = self.client.post(
                f'/specifications/{specification_id}/items',
                json={'componentId': '1', 'quantityPerDevice': quantity},
                headers=self.headers,
            )
            self.assertEqual(response.status_code, 201)

        response = self.client.post(
            f'/specifications/{specification_id}/delivery', headers=self.headers,
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json['items'], 1)
        delivery = self.client.get('/deliveries', headers=self.headers).json['data'][0]
        self.assertEqual(delivery['id'], response.json['id'])
        self.assertEqual(delivery['name'], 'Дозаказ — Плата управления')
        self.assertEqual(delivery['sourceFile'], f'Спецификация #{specification_id}')
        self.assertEqual(len(delivery['items']), 1)
        item = delivery['items'][0]
        self.assertEqual(item['sourceRow'], 1)
        self.assertEqual(item['cnt'], '1')
        self.assertEqual(item['cellnum'], 'A-01')

        confirmed = self.client.post(
            f"/deliveries/{delivery['id']}/confirm", headers=self.headers,
        )
        self.assertEqual(confirmed.json, {'created': 0, 'merged': 1, 'items': 1})
        self.assertEqual(self.rows()[0][9], 4)

    def test_expected_delivery_requires_existing_specification_with_shortage(self):
        created = self.client.post(
            '/specifications', json={'name': 'Без дефицита', 'deviceQuantity': 1},
            headers=self.headers,
        )
        specification_id = created.json['id']
        response = self.client.post(
            f'/specifications/{specification_id}/delivery', headers=self.headers,
        )
        self.assertEqual(response.status_code, 409)
        self.assertIn('больше не требуется', response.json['error'])
        self.assertEqual(self.client.post(
            '/specifications/999/delivery', headers=self.headers,
        ).status_code, 404)
        self.assertEqual(self.client.get('/deliveries', headers=self.headers).json['data'], [])

    def test_unmatched_specification_item_is_included_in_reorder_export(self):
        created = self.client.post(
            '/specifications', json={'name': 'Новая плата', 'deviceQuantity': '4'},
            headers=self.headers,
        )
        specification_id = created.json['id']
        item = {
            'componentId': '', 'group': 'Конденсатор', 'name': '',
            'value': '100', 'unit': 'нФ', 'tol': '10%',
            'description': 'Керамический\nX7R', 'case': '0603',
            'manufacturer': 'Murata', 'quantityPerDevice': '3',
        }
        self.assertEqual(self.client.post(
            f'/specifications/{specification_id}/items', json=item, headers=self.headers,
        ).status_code, 201)
        specification = self.client.get('/specifications', headers=self.headers).json['data'][0]
        self.assertEqual(specification['items'][0]['status'], 'unmatched')
        self.assertEqual(specification['summary']['toOrder'], 12)
        exported = self.client.get(
            f'/specifications/{specification_id}/export', headers=self.headers,
        )
        self.assertEqual(exported.status_code, 200)
        workbook = __import__('openpyxl').load_workbook(BytesIO(exported.data), data_only=True)
        self.assertEqual(list(workbook.active.values), [
            ('#', 'Тип элемента', 'Наименование', 'Описание', 'Количество'),
            (1, 'Конденсатор', '100 нФ 10% 0603 Murata', 'Керамический\nX7R', 12),
        ])
        workbook.close()

    def test_unmatched_specification_is_rematched_after_stock_appears(self):
        created = self.client.post(
            '/specifications', json={'name': 'Старая спецификация', 'deviceQuantity': '2'},
            headers=self.headers,
        )
        specification_id = created.json['id']
        item = {
            'componentId': '', 'group': 'Конденсатор', 'name': 'CC0603JRNPO**120',
            'value': '12', 'unit': 'пФ', 'tol': '5%', 'description': '',
            'case': '0603', 'manufacturer': 'Yageo', 'quantityPerDevice': '2',
        }
        self.client.post(
            f'/specifications/{specification_id}/items', json=item, headers=self.headers,
        )
        before = self.client.get('/specifications', headers=self.headers).json['data'][0]['items'][0]
        self.assertEqual(before['status'], 'unmatched')

        self.post(
            'Add', group='Конденсатор', name='CC0603JRNPO**120', value='12.0',
            unit='pF', tol='5 %', description='12pF 5% 50V NP0 0603',
            case='0603', manufacturer='Yageo', cnt='7', cellnum='',
        )
        # Старая ссылка может остаться после удаления и повторного добавления склада.
        self.db.SpecificationItem.update(ComponentID=999).execute()
        after = self.client.get('/specifications', headers=self.headers).json['data'][0]['items'][0]
        self.assertEqual(after['componentId'], '1')
        self.assertEqual(after['stockQuantity'], '7')
        self.assertEqual(after['status'], 'enough')
        stored = self.db.SpecificationItem.get_by_id(after['id'])
        self.assertEqual(stored.ComponentID, 1)

    def test_compact_bom_with_mojibake_headers_and_english_units_is_supported(self):
        self.post(
            'Add', group='Конденсатор', name='GRM188R71C104KA01D', value='0.1',
            unit='мкФ', tol='5%', description='100nF X7R', case='0603',
            manufacturer='Murata', cnt='20', cellnum='C-01',
        )
        response = self.client.post(
            '/specifications/import',
            data={
                'name': 'Компактный BOM', 'deviceQuantity': '3',
                'file': (self.compact_specification_report(), 'ResultTable (2).xlsx'),
            },
            content_type='multipart/form-data', headers=self.headers,
        )
        self.assertEqual(response.status_code, 201)
        specification = self.client.get('/specifications', headers=self.headers).json['data'][0]
        capacitor, resistor, microcontroller = specification['items']
        self.assertEqual(
            [capacitor['group'], capacitor['value'], capacitor['unit'], capacitor['componentId']],
            ['Конденсатор', '0.1', 'мкФ', '1'],
        )
        self.assertEqual(
            [resistor['group'], resistor['value'], resistor['unit'], resistor['status']],
            ['Резистор', '10', 'кОм', 'unmatched'],
        )
        self.assertEqual(microcontroller['group'], 'Прочее')


if __name__ == '__main__':
    unittest.main()
